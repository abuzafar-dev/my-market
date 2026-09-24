"""Sales report files for one period (day / week / month).

Two formats, both written for a shop owner rather than an accountant:
  * .xlsx — a workbook with four sheets: summary, per-day (or per-hour),
    every single sale, and per-product totals;
  * .csv  — just the per-day table, in a form Excel opens with the right
    columns and letters (BOM + ";" separator).

Headers and day names come in the language the app is currently showing
(`lang=uz|ru`).
"""

import csv
from datetime import date
from io import BytesIO

from django.db.models import Sum
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.sales.models import SaleItem
from apps.shops.models import Shop

from .services import completed_sales, debt_summary, sales_series

_XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

TEXT = {
    "uz": {
        "title": "Savdo hisoboti",
        "shop": "Do'kon",
        "period": "Davr",
        "from_to": "Sanalar",
        "periods": {"day": "Kun", "week": "Hafta", "month": "Oy"},
        "summary": "Xulosa",
        "days": "Kunlar bo'yicha",
        "hours": "Soatlar bo'yicha",
        "sales": "Barcha sotuvlar",
        "products": "Mahsulotlar bo'yicha",
        "date": "Sana",
        "hour": "Soat",
        "weekday": "Hafta kuni",
        "count": "Sotuvlar soni",
        "debt_paid": "Davrda qabul qilingan qarz to'lovlari (so'm)",
        "debt_outstanding": "Hozirgi jami qarz qoldig'i (so'm)",
        "revenue": "Jami savdo (so'm)",
        "cash": "Naqd (so'm)",
        "card": "Karta (so'm)",
        "debt": "Qarzga (so'm)",
        "profit": "Foyda (so'm)",
        "total": "JAMI",
        "datetime": "Sana va vaqt",
        "receipt": "Chek raqami",
        "seller": "Sotgan xodim",
        "customer": "Mijoz",
        "payment": "To'lov turi",
        "amount": "Summa (so'm)",
        "items": "Mahsulotlar",
        "product": "Mahsulot",
        "unit": "Birlik",
        "qty": "Sotilgan miqdor",
        "product_revenue": "Tushum (so'm)",
        "weekdays": [
            "Dushanba",
            "Seshanba",
            "Chorshanba",
            "Payshanba",
            "Juma",
            "Shanba",
            "Yakshanba",
        ],
        "payment_types": {"cash": "Naqd", "card": "Karta", "debt": "Qarz"},
        "units": {"kg": "kg", "liter": "litr", "piece": "dona"},
        "file": "hisobot",
    },
    "ru": {
        "title": "Отчёт о продажах",
        "shop": "Магазин",
        "period": "Период",
        "from_to": "Даты",
        "periods": {"day": "День", "week": "Неделя", "month": "Месяц"},
        "summary": "Итоги",
        "days": "По дням",
        "hours": "По часам",
        "sales": "Все продажи",
        "products": "По товарам",
        "date": "Дата",
        "hour": "Час",
        "weekday": "День недели",
        "count": "Кол-во продаж",
        "debt_paid": "Оплаты долгов за период (сум)",
        "debt_outstanding": "Общий долг клиентов сейчас (сум)",
        "revenue": "Выручка (сум)",
        "cash": "Наличные (сум)",
        "card": "Карта (сум)",
        "debt": "В долг (сум)",
        "profit": "Прибыль (сум)",
        "total": "ИТОГО",
        "datetime": "Дата и время",
        "receipt": "Номер чека",
        "seller": "Продавец",
        "customer": "Клиент",
        "payment": "Вид оплаты",
        "amount": "Сумма (сум)",
        "items": "Товары",
        "product": "Товар",
        "unit": "Единица",
        "qty": "Продано",
        "product_revenue": "Выручка (сум)",
        "weekdays": [
            "Понедельник",
            "Вторник",
            "Среда",
            "Четверг",
            "Пятница",
            "Суббота",
            "Воскресенье",
        ],
        "payment_types": {"cash": "Наличные", "card": "Карта", "debt": "В долг"},
        "units": {"kg": "кг", "liter": "л", "piece": "шт"},
        "file": "otchet",
    },
}

_FORMULA_STARTS = ("=", "+", "-", "@", "\t", "\r")


def safe_cell(value):
    """Spreadsheet formula injection guard. A product / customer / staff name is
    free text, and Excel or Sheets run a cell that starts with = + - @ as a
    formula (``=HYPERLINK(...)``, DDE...) the moment the owner opens the file.
    A leading apostrophe makes it plain text. Numbers pass through untouched."""
    if isinstance(value, str) and value.startswith(_FORMULA_STARTS):
        return "'" + value
    return value


def safe_row(values) -> list:
    return [safe_cell(value) for value in values]


_MONEY_FIELDS = ("revenue", "cash", "card", "debt", "profit")
_MONEY_FORMAT = "#,##0"
_HEADER_FILL = PatternFill("solid", fgColor="0F172A")
_TOTAL_FILL = PatternFill("solid", fgColor="F0FDFA")


def _filename(lang: str, start: date, end: date, ext: str) -> str:
    span = start.isoformat() if start == end else f"{start.isoformat()}_{end.isoformat()}"
    return f"{TEXT[lang]['file']}_{span}.{ext}"


def _day_table(shop: Shop, start: date, end: date, lang: str) -> tuple[list[str], list[list], list]:
    """Header, body rows and the totals row of the per-day (per-hour) table."""
    text = TEXT[lang]
    series = sales_series(shop, start, end)
    is_hourly = series["kind"] == "hour"

    header = [text["hour"] if is_hourly else text["date"]]
    if not is_hourly:
        header.append(text["weekday"])
    header += [text["count"], text["revenue"], text["cash"], text["card"], text["debt"]]
    header.append(text["profit"])

    body = []
    totals = {"count": 0, **{field: 0 for field in _MONEY_FIELDS}}
    for row in series["rows"]:
        if is_hourly:
            lead = [row["key"]]
        else:
            day = date.fromisoformat(row["key"])
            lead = [day.strftime("%d.%m.%Y"), text["weekdays"][day.weekday()]]
        body.append(
            [*lead, row["count"], *(row[field] for field in _MONEY_FIELDS)],
        )
        totals["count"] += row["count"]
        for field in _MONEY_FIELDS:
            totals[field] += row[field]

    blanks = [""] if not is_hourly else []
    total_row = [text["total"], *blanks, totals["count"], *(totals[f] for f in _MONEY_FIELDS)]
    return header, body, total_row


def period_csv_response(shop: Shop, start: date, end: date, lang: str) -> HttpResponse:
    header, body, total_row = _day_table(shop, start, end, lang)

    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{_filename(lang, start, end, "csv")}"'
    # BOM so Excel reads UTF-8 (Uzbek / Russian letters); ";" because Excel in
    # uz/ru regional settings splits columns on ";" and would leave a "," file
    # in a single column.
    response.write("\ufeff")
    writer = csv.writer(response, delimiter=";")
    writer.writerow(header)
    writer.writerows(body)
    writer.writerow(total_row)
    return response


def _style_header(sheet, row: int = 1) -> None:
    for cell in sheet[row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[row].height = 30


def _format_money_columns(sheet, header: list[str], money_headers: set[str]) -> None:
    for index, title in enumerate(header, start=1):
        if title in money_headers:
            for cell in sheet[get_column_letter(index)][1:]:
                cell.number_format = _MONEY_FORMAT


def _autosize(sheet) -> None:
    for column in sheet.columns:
        width = max((len(str(cell.value)) for cell in column if cell.value is not None), default=0)
        sheet.column_dimensions[column[0].column_letter].width = min(max(width + 3, 12), 48)


def period_xlsx_response(
    shop: Shop, period: str, start: date, end: date, lang: str
) -> HttpResponse:
    text = TEXT[lang]
    header, body, total_row = _day_table(shop, start, end, lang)
    money_headers = {text[f] for f in (*_MONEY_FIELDS, "debt_paid", "debt_outstanding")}

    workbook = Workbook()

    # 1) Summary — the first thing the owner sees when the file opens.
    summary = workbook.active
    summary.title = text["summary"]
    summary["A1"] = text["title"]
    summary["A1"].font = Font(bold=True, size=16)
    summary.append([])
    summary.append(safe_row([text["shop"], shop.name]))
    summary.append([text["period"], text["periods"][period]])
    dates = start.strftime("%d.%m.%Y")
    if start != end:
        dates += " — " + end.strftime("%d.%m.%Y")
    summary.append([text["from_to"], dates])
    summary.append([])
    labels = [text["count"], *(text[f] for f in _MONEY_FIELDS)]
    total_by_label = dict(zip(header[-6:], total_row[-6:], strict=True))
    for label in labels:
        summary.append([label, total_by_label[label]])
    debt = debt_summary(shop, start, end)
    summary.append([])
    summary.append([text["debt_paid"], debt["paid"]])
    summary.append([text["debt_outstanding"], debt["outstanding"]])
    for row in summary.iter_rows(min_row=3, max_row=summary.max_row):
        row[0].font = Font(bold=True)
        if row[0].value in money_headers:
            row[1].number_format = _MONEY_FORMAT
        row[1].alignment = Alignment(horizontal="left")
    summary.column_dimensions["A"].width = 44
    summary.column_dimensions["B"].width = 30

    # 2) Per day / per hour
    kind_sheet = workbook.create_sheet(text["hours"] if start == end else text["days"])
    kind_sheet.append(header)
    for row in body:
        kind_sheet.append(row)
    kind_sheet.append(total_row)
    _style_header(kind_sheet)
    for cell in kind_sheet[kind_sheet.max_row]:
        cell.font = Font(bold=True)
        cell.fill = _TOTAL_FILL
    _format_money_columns(kind_sheet, header, money_headers)
    kind_sheet.freeze_panes = "A2"
    _autosize(kind_sheet)

    # 3) Every sale
    sales_sheet = workbook.create_sheet(text["sales"])
    sales_header = [
        text["datetime"],
        text["receipt"],
        text["seller"],
        text["customer"],
        text["payment"],
        text["amount"],
        text["profit"],
        text["items"],
    ]
    sales_sheet.append(sales_header)
    sales = (
        completed_sales(shop, start, end)
        .select_related("sold_by", "customer")
        .prefetch_related("items__product")
        .order_by("sold_at")
    )
    for sale in sales:
        sales_sheet.append(
            safe_row(
                [
                    timezone.localtime(sale.sold_at).strftime("%d.%m.%Y %H:%M"),
                    str(sale.id)[:8].upper(),
                    sale.sold_by.full_name if sale.sold_by else "—",
                    sale.customer.full_name if sale.customer else "",
                    text["payment_types"].get(sale.payment_type, sale.payment_type),
                    sale.total,
                    sale.total - sale.cost_total,
                    ", ".join(
                        f"{item.product.name} × {item.qty.normalize():f}"
                        for item in sale.items.all()
                    ),
                ]
            )
        )
    _style_header(sales_sheet)
    _format_money_columns(sales_sheet, sales_header, {text["amount"], text["profit"]})
    sales_sheet.freeze_panes = "A2"
    _autosize(sales_sheet)

    # 4) Per product
    products_sheet = workbook.create_sheet(text["products"])
    products_header = [text["product"], text["unit"], text["qty"], text["product_revenue"]]
    products_sheet.append(products_header)
    per_product = (
        SaleItem.objects.filter(
            sale__in=completed_sales(shop, start, end),
        )
        .values("product__name", "product__unit")
        .annotate(qty_sold=Sum("qty"), revenue=Sum("line_total"))
        .order_by("-revenue")
    )
    for row in per_product:
        products_sheet.append(
            safe_row(
                [
                    row["product__name"],
                    text["units"].get(row["product__unit"], row["product__unit"]),
                    float(row["qty_sold"]),
                    row["revenue"],
                ]
            )
        )
    _style_header(products_sheet)
    _format_money_columns(products_sheet, products_header, {text["product_revenue"]})
    products_sheet.freeze_panes = "A2"
    _autosize(products_sheet)

    buffer = BytesIO()
    workbook.save(buffer)
    response = HttpResponse(buffer.getvalue(), content_type=_XLSX_CONTENT_TYPE)
    response["Content-Disposition"] = (
        f'attachment; filename="{_filename(lang, start, end, "xlsx")}"'
    )
    return response
