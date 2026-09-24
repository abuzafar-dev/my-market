from datetime import date
from decimal import Decimal
from io import BytesIO
from uuid import uuid4

from django.test import TestCase
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.catalog.models import Batch, Product
from apps.sales.models import Sale
from apps.sales.services import CartLine, cancel_sale, create_sale
from apps.shops.models import Shop, ShopSettings, User


class ReportExportTests(TestCase):
    def setUp(self):
        self.shop = Shop.objects.create(name="Do'kon", phone="+998900000000")
        ShopSettings.objects.create(shop=self.shop)
        self.owner = User.objects.create_user(
            phone="+998901112233", password="pass1234", shop=self.shop, full_name="Egasi"
        )
        self.seller = User.objects.create_user(
            phone="+998907778899",
            password="pass1234",
            shop=self.shop,
            full_name="Sotuvchi",
            role=User.Role.SELLER,
        )
        self.client = APIClient()
        self.client.force_authenticate(self.owner)

    def make_product(self, name, stock, min_stock="5", unit=Product.Unit.KG):
        product = Product.objects.create(
            shop=self.shop,
            name=name,
            unit=unit,
            markup_amount=2000,
            min_stock=Decimal(min_stock),
        )
        Batch.objects.create(
            shop=self.shop,
            product=product,
            qty_initial=Decimal(stock),
            qty_remaining=Decimal(stock),
            cost_price=1000,
            sale_price=1200,
        )
        return product

    def sell(self, product, qty):
        return create_sale(
            shop=self.shop,
            user=self.owner,
            client_id=uuid4(),
            payment_type=Sale.PaymentType.CASH,
            customer=None,
            cart=[CartLine(product=product, qty=Decimal(qty))],
        )

    def sheet_rows(self, response):
        sheet = load_workbook(BytesIO(response.content)).active
        return [[cell.value for cell in row] for row in sheet.iter_rows()]

    def test_low_stock_export_lists_only_low_products(self):
        self.make_product("Un", "3")
        self.make_product("Shakar", "50")

        response = self.client.get("/api/reports/low-stock/export/")

        self.assertEqual(response.status_code, 200)
        rows = self.sheet_rows(response)
        self.assertEqual(rows[0][0], "Nomi")
        self.assertEqual([row[0] for row in rows[1:]], ["Un"])
        self.assertEqual(rows[1][3], 3.0)

    def test_unsold_export_excludes_products_with_completed_sales(self):
        sold = self.make_product("Un", "50")
        self.make_product("Guruch", "50")
        self.sell(sold, "2")

        response = self.client.get("/api/reports/unsold/export/")

        self.assertEqual(response.status_code, 200)
        self.assertEqual([row[0] for row in self.sheet_rows(response)[1:]], ["Guruch"])

    def test_unsold_export_counts_cancelled_sale_as_unsold(self):
        product = self.make_product("Un", "50")
        cancel_sale(self.sell(product, "2"), self.owner)

        response = self.client.get("/api/reports/unsold/export/")

        self.assertEqual([row[0] for row in self.sheet_rows(response)[1:]], ["Un"])

    def test_exports_are_owner_only(self):
        self.client.force_authenticate(self.seller)
        for url in ("/api/reports/low-stock/export/", "/api/reports/unsold/export/"):
            self.assertEqual(self.client.get(url).status_code, 403)

    def test_sold_products_list_everything_with_qty_profit_and_receipts(self):
        flour, bread = self.make_product("Un", "50"), self.make_product("Non", "50")
        self.sell(flour, "2.5")  # 3000, profit 500
        self.sell(flour, "1")  # 1200, profit 200
        self.sell(bread, "1")  # 1200, profit 200

        response = self.client.get("/api/reports/", {"period": "day"})

        rows = response.json()["data"]["products"]
        self.assertEqual([row["name"] for row in rows], ["Un", "Non"])
        self.assertEqual(rows[0]["unit"], "kg")
        self.assertEqual(Decimal(rows[0]["qty"]), Decimal("3.5"))
        self.assertEqual((rows[0]["revenue"], rows[0]["profit"]), (4200, 700))
        self.assertEqual(rows[0]["receipts"], 2)

    def test_any_past_day_can_be_opened(self):
        from datetime import timedelta

        from django.utils import timezone

        product = self.make_product("Un", "50")
        sale = self.sell(product, "2")
        three_days_ago = timezone.now() - timedelta(days=3)
        Sale.objects.filter(pk=sale.pk).update(sold_at=three_days_ago)
        day = timezone.localdate(three_days_ago).isoformat()

        data = self.client.get("/api/reports/", {"period": "day", "date": day}).json()["data"]

        self.assertEqual(data["range"], {"start": day, "end": day})
        self.assertEqual(data["stats"]["revenue"], 2400)
        self.assertEqual(data["products"][0]["name"], "Un")
        today = self.client.get("/api/reports/", {"period": "day"}).json()["data"]
        self.assertEqual(today["stats"]["revenue"], 0)

    def test_a_bad_date_is_a_validation_error_and_a_future_one_means_today(self):
        from django.utils import timezone

        bad = self.client.get("/api/reports/", {"date": "24.09.2026"})
        self.assertEqual(bad.status_code, 400)

        future = self.client.get("/api/reports/", {"date": "2999-01-01"}).json()["data"]
        today = timezone.localdate().isoformat()
        self.assertEqual(future["range"], {"start": today, "end": today})
        self.assertEqual(future["today"], today)


class PeriodReportTests(ReportExportTests):
    """Per-day breakdown and the period export files."""

    def test_month_series_lists_every_day_up_to_today_and_puts_sales_on_today(self):
        self.sell(self.make_product("Un", "50"), "2")  # revenue 2400, profit 400

        data = self.client.get("/api/reports/", {"period": "month"}).json()["data"]

        rows = data["series"]["rows"]
        self.assertEqual(data["series"]["kind"], "day")
        self.assertEqual(len(rows), date.today().day)  # 1st .. today, zeros included
        self.assertEqual(rows[0]["key"], date.today().replace(day=1).isoformat())
        self.assertEqual(rows[-1]["key"], date.today().isoformat())
        self.assertEqual(
            (rows[-1]["count"], rows[-1]["revenue"], rows[-1]["profit"]), (1, 2400, 400)
        )
        self.assertEqual(sum(row["revenue"] for row in rows), data["stats"]["revenue"])
        self.assertEqual(data["range"]["end"], date.today().isoformat())

    def test_week_series_starts_on_monday(self):
        rows = self.client.get("/api/reports/", {"period": "week"}).json()["data"]["series"]["rows"]

        first = date.fromisoformat(rows[0]["key"])
        self.assertEqual(first.weekday(), 0)
        self.assertEqual(len(rows), date.today().weekday() + 1)

    def test_day_series_is_by_hour_and_only_hours_with_sales(self):
        self.sell(self.make_product("Un", "50"), "1")

        series = self.client.get("/api/reports/", {"period": "day"}).json()["data"]["series"]

        self.assertEqual(series["kind"], "hour")
        self.assertEqual(len(series["rows"]), 1)
        self.assertRegex(series["rows"][0]["key"], r"^\d\d:00$")

    def test_cancelled_sales_are_not_in_the_series(self):
        sale = self.sell(self.make_product("Un", "50"), "1")
        cancel_sale(sale, self.owner)

        rows = self.client.get("/api/reports/", {"period": "week"}).json()["data"]["series"]["rows"]

        self.assertEqual(sum(row["revenue"] for row in rows), 0)

    def test_xlsx_export_has_summary_days_sales_and_products_sheets(self):
        product = self.make_product("Un", "50")
        self.sell(product, "2")
        self.sell(product, "1")

        response = self.client.get(
            "/api/reports/export/", {"period": "week", "file": "xlsx", "lang": "uz"}
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn(".xlsx", response["Content-Disposition"])
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(
            workbook.sheetnames,
            ["Xulosa", "Kunlar bo'yicha", "Barcha sotuvlar", "Mahsulotlar bo'yicha"],
        )
        days = [[c.value for c in row] for row in workbook["Kunlar bo'yicha"].iter_rows()]
        self.assertEqual(days[0][:3], ["Sana", "Hafta kuni", "Sotuvlar soni"])
        self.assertEqual(days[-1][0], "JAMI")
        self.assertEqual(days[-1][2:4], [2, 3600])  # 2 sales, 3 kg * 1200
        sales = list(workbook["Barcha sotuvlar"].iter_rows(values_only=True))
        self.assertEqual(len(sales), 3)  # header + 2 sales
        products = list(workbook["Mahsulotlar bo'yicha"].iter_rows(values_only=True))
        self.assertEqual(products[1], ("Un", "kg", 3.0, 3600))

    def test_xlsx_export_in_russian(self):
        self.sell(self.make_product("Un", "50"), "1")

        response = self.client.get(
            "/api/reports/export/", {"period": "month", "file": "xlsx", "lang": "ru"}
        )

        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(workbook.sheetnames[0], "Итоги")
        self.assertIn("otchet_", response["Content-Disposition"])

    def test_csv_export_is_excel_friendly_and_matches_the_screen_totals(self):
        self.sell(self.make_product("Un", "50"), "2")

        response = self.client.get("/api/reports/export/", {"period": "month", "file": "csv"})

        self.assertEqual(response.status_code, 200)
        body = response.content.decode("utf-8")
        self.assertTrue(body.startswith("﻿"))  # BOM: Excel reads the letters right
        lines = body.lstrip("﻿").splitlines()
        self.assertTrue(lines[0].startswith("Sana;Hafta kuni;Sotuvlar soni"))
        self.assertTrue(lines[-1].startswith("JAMI;;1;2400"))

    def test_export_is_owner_only(self):
        seller = APIClient()
        seller.force_authenticate(self.seller)

        self.assertEqual(seller.get("/api/reports/export/", {"file": "xlsx"}).status_code, 403)


class DebtSummaryTests(ReportExportTests):
    """The credit block of the report: sold on credit, paid back, owed."""

    def make_customer(self, name):
        from apps.debt.models import Customer

        return Customer.objects.create(shop=self.shop, full_name=name, phone="+998901112233")

    def sell_on_credit(self, product, qty, customer):
        return create_sale(
            shop=self.shop,
            user=self.owner,
            client_id=uuid4(),
            payment_type=Sale.PaymentType.DEBT,
            customer=customer,
            cart=[CartLine(product=product, qty=Decimal(qty))],
        )

    def debt(self, period="month"):
        return self.client.get("/api/reports/", {"period": period}).json()["data"]["debt"]

    def test_credit_sales_are_summed_and_counted_apart_from_cash(self):
        product = self.make_product("Un", "100")
        aziz, vali = self.make_customer("Aziz"), self.make_customer("Vali")
        self.sell_on_credit(product, "2", aziz)  # 2400
        self.sell_on_credit(product, "1", aziz)  # 1200
        self.sell_on_credit(product, "5", vali)  # 6000
        self.sell(product, "10")  # cash — must not count

        debt = self.debt()

        self.assertEqual((debt["sold"], debt["sold_count"], debt["sold_customers"]), (9600, 3, 2))

    def test_payments_and_outstanding_balance(self):
        from apps.debt.services import add_payment

        product = self.make_product("Un", "100")
        aziz = self.make_customer("Aziz")
        self.sell_on_credit(product, "5", aziz)  # owes 6000
        add_payment(customer=aziz, user=self.owner, amount=2500)

        debt = self.debt()

        self.assertEqual(debt["paid"], 2500)
        self.assertEqual((debt["outstanding"], debt["debtors"]), (3500, 1))

    def test_cancelled_credit_sale_is_not_counted(self):
        product = self.make_product("Un", "100")
        sale = self.sell_on_credit(product, "2", self.make_customer("Aziz"))
        cancel_sale(sale, self.owner)

        self.assertEqual(self.debt()["sold"], 0)

    def test_cancelled_credit_sale_is_not_a_repayment(self):
        # Cancelling books a reversing PAYMENT entry; no money came in.
        product = self.make_product("Un", "100")
        sale = self.sell_on_credit(product, "2", self.make_customer("Aziz"))
        cancel_sale(sale, self.owner)

        self.assertEqual(self.debt()["paid"], 0)

    def test_dashboard_total_debt_ignores_overpaid_customers(self):
        from apps.debt.services import add_payment

        product = self.make_product("Un", "100")
        self.sell_on_credit(product, "5", self.make_customer("Aziz"))  # owes 6000
        add_payment(customer=self.make_customer("Vali"), user=self.owner, amount=1000)  # -1000

        data = self.client.get("/api/dashboard/").json()["data"]

        self.assertEqual(data["total_debt"], 6000)

    def test_xlsx_summary_carries_paid_and_outstanding(self):
        product = self.make_product("Un", "100")
        self.sell_on_credit(product, "5", self.make_customer("Aziz"))

        response = self.client.get(
            "/api/reports/export/", {"period": "month", "file": "xlsx", "lang": "uz"}
        )

        summary = {
            row[0]: row[1]
            for row in load_workbook(BytesIO(response.content))["Xulosa"].iter_rows(
                values_only=True
            )
            if row[0]
        }
        self.assertEqual(summary["Qarzga (so'm)"], 6000)
        self.assertEqual(summary["Hozirgi jami qarz qoldig'i (so'm)"], 6000)


class PreviousPeriodBoundsTests(TestCase):
    """The comparison window is the same stretch of the previous period."""

    def test_day_compares_with_yesterday(self):
        from apps.reports.services import previous_period_bounds

        today = date(2026, 9, 24)
        self.assertEqual(previous_period_bounds("day", today), (date(2026, 9, 23),) * 2)

    def test_week_compares_monday_to_same_weekday(self):
        from apps.reports.services import previous_period_bounds

        # 2026-09-24 is a Thursday: last week's Monday..Thursday.
        self.assertEqual(
            previous_period_bounds("week", date(2026, 9, 24)),
            (date(2026, 9, 14), date(2026, 9, 17)),
        )

    def test_month_clamps_to_shorter_previous_month(self):
        from apps.reports.services import previous_period_bounds

        self.assertEqual(
            previous_period_bounds("month", date(2026, 3, 31)),
            (date(2026, 2, 1), date(2026, 2, 28)),
        )
        self.assertEqual(
            previous_period_bounds("month", date(2026, 1, 15)),
            (date(2025, 12, 1), date(2025, 12, 15)),
        )


class LocalDateTests(ReportExportTests):
    """The shop's "today" is its own date (Asia/Tashkent), not the UTC server clock's."""

    def test_period_bounds_follow_shop_timezone(self):
        from datetime import datetime
        from unittest import mock
        from zoneinfo import ZoneInfo

        from apps.reports.services import period_bounds

        # 02:00 in Tashkent on 1 Jan is still 31 Dec in UTC.
        moment = datetime(2030, 1, 1, 2, 0, tzinfo=ZoneInfo("Asia/Tashkent"))
        with mock.patch("django.utils.timezone.now", return_value=moment):
            self.assertEqual(period_bounds("day"), (date(2030, 1, 1),) * 2)


class PeriodBoundsTests(TestCase):
    """The day / week / month around any date, never past today."""

    def bounds(self, period, anchor, today):
        from unittest import mock

        from apps.reports.services import period_bounds

        with mock.patch("django.utils.timezone.localdate", return_value=today):
            return period_bounds(period, anchor)

    def test_a_past_week_and_month_are_whole(self):
        today = date(2026, 9, 24)
        # 2026-08-12 is a Wednesday.
        self.assertEqual(
            self.bounds("week", date(2026, 8, 12), today), (date(2026, 8, 10), date(2026, 8, 16))
        )
        self.assertEqual(
            self.bounds("month", date(2026, 2, 10), today), (date(2026, 2, 1), date(2026, 2, 28))
        )

    def test_the_current_period_stops_at_today(self):
        today = date(2026, 9, 24)
        self.assertEqual(self.bounds("month", today, today), (date(2026, 9, 1), today))
        self.assertEqual(self.bounds("week", today, today), (date(2026, 9, 21), today))
