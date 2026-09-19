"""Security regression tests: what an attacker (or a curious neighbour shop)
must NOT be able to do. Each test names the attack it stops."""

import os
import subprocess
import sys
import tempfile
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch
from uuid import uuid4

from django.core.cache import cache
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, override_settings
from openpyxl import load_workbook
from PIL import Image
from rest_framework.test import APIClient, APITestCase
from rest_framework.throttling import UserRateThrottle

from apps.catalog.models import Batch, Category, Product
from apps.debt.models import Customer
from apps.sales.models import Sale
from apps.sales.services import CartLine, create_sale
from apps.shops.models import Shop, ShopSettings, User


def make_shop(name, phone):
    shop = Shop.objects.create(name=name, phone="+998900000000")
    ShopSettings.objects.create(shop=shop)
    owner = User.objects.create_user(
        phone=phone, password="Str0ng-pass!", shop=shop, full_name=name
    )
    return shop, owner


class TwoShopsTestCase(APITestCase):
    """Shop A is the attacker's; shop B holds the data that must stay private."""

    def setUp(self):
        cache.clear()
        self.addCleanup(cache.clear)
        self.shop_a, self.owner_a = make_shop("A", "+998901110001")
        self.shop_b, self.owner_b = make_shop("B", "+998901110002")

        self.cat_b = Category.objects.create(shop=self.shop_b, name="B-kategoriya")
        self.product_b = Product.objects.create(
            shop=self.shop_b,
            category=self.cat_b,
            name="B mahsuloti",
            barcode="4780000000999",
            unit=Product.Unit.PIECE,
            markup_amount=2000,
            min_stock=Decimal("1"),
        )
        self.batch_b = Batch.objects.create(
            shop=self.shop_b,
            product=self.product_b,
            qty_initial=Decimal("10"),
            qty_remaining=Decimal("10"),
            cost_price=1000,
            sale_price=1200,
        )
        self.customer_b = Customer.objects.create(shop=self.shop_b, full_name="B mijozi")
        self.sale_b = create_sale(
            shop=self.shop_b,
            user=self.owner_b,
            client_id=uuid4(),
            payment_type=Sale.PaymentType.CASH,
            customer=None,
            cart=[CartLine(product=self.product_b, qty=Decimal("1"))],
        )

        self.client = APIClient()
        self.client.force_authenticate(self.owner_a)


class TenantIsolationTests(TwoShopsTestCase):
    """IDOR: knowing another shop's ids must give an attacker nothing."""

    def assert_not_found(self, response):
        self.assertEqual(response.status_code, 404, response.content)

    def test_cannot_read_or_edit_another_shops_product(self):
        url = f"/api/products/{self.product_b.id}/"
        self.assert_not_found(self.client.get(url))
        self.assert_not_found(self.client.patch(url, {"name": "hacked"}, format="json"))
        self.assert_not_found(self.client.post(url + "archive/"))
        self.product_b.refresh_from_db()
        self.assertEqual((self.product_b.name, self.product_b.is_active), ("B mahsuloti", True))

    def test_cannot_look_up_another_shops_barcode(self):
        self.assert_not_found(self.client.get(f"/api/products/barcode/{self.product_b.barcode}/"))

    def test_cannot_stock_or_write_off_another_shops_goods(self):
        stock_in = self.client.post(
            "/api/batches/",
            {"product": str(self.product_b.id), "qty_initial": 5, "cost_price": 1000},
            format="json",
        )
        self.assertEqual(stock_in.status_code, 400)
        self.assert_not_found(
            self.client.post(
                f"/api/batches/{self.batch_b.id}/writeoff/",
                {"qty": 1, "reason": "lost"},
                format="json",
            )
        )
        self.batch_b.refresh_from_db()
        self.assertEqual(self.batch_b.qty_remaining, Decimal("9"))  # only setUp's own sale

    def test_cannot_touch_another_shops_customer(self):
        base = f"/api/customers/{self.customer_b.id}/"
        self.assert_not_found(self.client.get(base))
        self.assert_not_found(self.client.patch(base, {"full_name": "x"}, format="json"))
        self.assert_not_found(self.client.post(base + "debt/", {"amount": 1000}, format="json"))
        self.assert_not_found(self.client.post(base + "payment/", {"amount": 1000}, format="json"))
        self.customer_b.refresh_from_db()
        self.assertEqual(self.customer_b.debt_balance, 0)

    def test_cannot_read_or_cancel_another_shops_sale(self):
        self.assert_not_found(self.client.get(f"/api/sales/{self.sale_b.id}/"))
        self.assert_not_found(self.client.post(f"/api/sales/{self.sale_b.id}/cancel/"))
        self.sale_b.refresh_from_db()
        self.assertEqual(self.sale_b.status, Sale.Status.COMPLETED)

    def test_cannot_sell_another_shops_product_or_bill_its_customer(self):
        item = [{"product_id": str(self.product_b.id), "qty": 1}]
        sale = self.client.post(
            "/api/sales/",
            {"client_id": str(uuid4()), "payment_type": "cash", "items": item},
            format="json",
        )
        self.assertEqual(sale.status_code, 400)

        own = Product.objects.create(
            shop=self.shop_a,
            name="A",
            unit=Product.Unit.PIECE,
            markup_amount=1000,
            min_stock=Decimal("0"),
        )
        Batch.objects.create(
            shop=self.shop_a,
            product=own,
            qty_initial=1,
            qty_remaining=1,
            cost_price=100,
            sale_price=110,
        )
        on_credit = self.client.post(
            "/api/sales/",
            {
                "client_id": str(uuid4()),
                "payment_type": "debt",
                "customer_id": str(self.customer_b.id),
                "items": [{"product_id": str(own.id), "qty": 1}],
            },
            format="json",
        )
        self.assert_not_found(on_credit)

    def test_cannot_attach_a_product_to_another_shops_category(self):
        response = self.client.post(
            "/api/products/",
            {
                "name": "x",
                "category": str(self.cat_b.id),
                "unit": "piece",
                "markup_amount": 1000,
                "min_stock": 0,
            },
            format="json",
        )
        self.assertEqual(response.status_code, 400)

    def test_lists_and_reports_only_show_own_shop(self):
        for path in ("/api/products/", "/api/customers/", "/api/sales/"):
            body = self.client.get(path).json()["data"]
            self.assertEqual(len(body["results"]), 0, path)
        report = self.client.get("/api/reports/", {"period": "month"}).json()["data"]
        self.assertEqual(report["stats"]["revenue"], 0)
        self.assertEqual(report["top_products"], [])
        self.assertEqual(self.client.get("/api/dashboard/").json()["data"]["total_debt"], 0)


class AuthenticationRequiredTests(APITestCase):
    def test_every_api_endpoint_rejects_anonymous_requests(self):
        anyid = "00000000-0000-0000-0000-000000000000"
        for method, path in [
            ("get", "/api/products/"),
            ("post", "/api/products/"),
            ("get", f"/api/products/{anyid}/"),
            ("get", "/api/products/barcode/123/"),
            ("get", "/api/categories/"),
            ("post", "/api/batches/"),
            ("post", f"/api/batches/{anyid}/writeoff/"),
            ("get", "/api/purchase-list/"),
            ("get", "/api/customers/"),
            ("post", f"/api/customers/{anyid}/debt/"),
            ("get", "/api/sales/"),
            ("post", "/api/sales/"),
            ("post", f"/api/sales/{anyid}/cancel/"),
            ("get", "/api/dashboard/"),
            ("get", "/api/reports/"),
            ("get", "/api/reports/export/"),
            ("get", "/api/reports/low-stock/export/"),
            ("get", "/api/reports/unsold/export/"),
            ("get", "/api/settings/"),
            ("post", "/api/auth/password/"),
            ("post", "/api/auth/logout/"),
        ]:
            with self.subTest(path=path):
                response = getattr(self.client, method)(path)
                self.assertEqual(response.status_code, 401, f"{method} {path}")

    def test_a_forged_or_garbage_token_is_rejected(self):
        for token in ("garbage", "a.b.c", ""):
            response = self.client.get("/api/products/", HTTP_AUTHORIZATION=f"Bearer {token}")
            self.assertEqual(response.status_code, 401)


class InputLimitTests(TwoShopsTestCase):
    """Absurd numbers/lengths must be a clean 400, never a 500 or a bloated row."""

    def test_huge_debt_amount_and_long_note_are_rejected(self):
        own = Customer.objects.create(shop=self.shop_a, full_name="Mijoz")
        url = f"/api/customers/{own.id}/debt/"
        self.assertEqual(self.client.post(url, {"amount": 10**30}, format="json").status_code, 400)
        self.assertEqual(
            self.client.post(
                url, {"amount": 1000, "note": "x" * 10_000}, format="json"
            ).status_code,
            400,
        )
        self.assertEqual(self.client.post(url, {"amount": 1000}, format="json").status_code, 200)

    def test_huge_quantity_and_too_many_lines_are_rejected(self):
        product = str(self.product_b.id)
        body = {"client_id": str(uuid4()), "payment_type": "cash"}
        huge_qty = {**body, "items": [{"product_id": product, "qty": "99999999999"}]}
        many = {**body, "items": [{"product_id": product, "qty": 1}] * 101}
        self.assertEqual(self.client.post("/api/sales/", huge_qty, format="json").status_code, 400)
        self.assertEqual(self.client.post("/api/sales/", many, format="json").status_code, 400)

    def test_huge_prices_are_rejected(self):
        own = Product.objects.create(
            shop=self.shop_a,
            name="A",
            unit=Product.Unit.PIECE,
            markup_amount=1000,
            min_stock=Decimal("0"),
        )
        for payload in (
            {"cost_price": 10**15},
            {"cost_price": 1000, "sale_price": 10**15},
            {"cost_price": 10**10},  # in range itself, but cost + markup is not
        ):
            response = self.client.post(
                "/api/batches/",
                {"product": str(own.id), "qty_initial": 1, **payload},
                format="json",
            )
            self.assertEqual(response.status_code, 400, payload)

    def test_client_cannot_set_server_owned_fields(self):
        created = self.client.post(
            "/api/customers/",
            {"full_name": "Yangi", "debt_balance": 999999, "is_active": False},
            format="json",
        )
        self.assertEqual(created.status_code, 201)
        customer = Customer.objects.get(pk=created.json()["data"]["id"])
        self.assertEqual((customer.debt_balance, customer.is_active), (0, True))

        own = Product.objects.create(
            shop=self.shop_a,
            name="A",
            unit=Product.Unit.PIECE,
            markup_amount=1000,
            min_stock=Decimal("0"),
        )
        self.client.patch(
            f"/api/products/{own.id}/",
            {"is_active": False, "shop": str(self.shop_b.id)},
            format="json",
        )
        own.refresh_from_db()
        self.assertEqual((own.is_active, own.shop_id), (True, self.shop_a.id))

    def test_absurdly_long_login_fields_are_rejected_before_hashing(self):
        response = APIClient().post(
            "/api/auth/login/", {"phone": "1" * 10_000, "password": "x" * 1_000_000}, format="json"
        )
        self.assertEqual(response.status_code, 400)


class LoginLockoutTests(TwoShopsTestCase):
    def login(self, password, ip="10.0.0.1", phone="+998901110001"):
        return APIClient(REMOTE_ADDR=ip).post(
            "/api/auth/login/", {"phone": phone, "password": password}, format="json"
        )

    def test_five_wrong_passwords_lock_out_even_the_right_one(self):
        for _ in range(5):
            self.assertEqual(self.login("wrong").status_code, 400)

        locked = self.login("Str0ng-pass!")

        self.assertEqual(locked.status_code, 429)

    def test_lockout_is_per_address_so_the_real_owner_can_still_log_in(self):
        for _ in range(5):
            self.login("wrong", ip="6.6.6.6")

        self.assertEqual(self.login("Str0ng-pass!", ip="6.6.6.6").status_code, 429)
        self.assertEqual(self.login("Str0ng-pass!", ip="10.0.0.9").status_code, 200)

    def test_a_number_attacked_from_many_addresses_is_locked_for_everyone(self):
        for n in range(30):
            self.login("wrong", ip=f"7.7.7.{n}")

        self.assertEqual(self.login("Str0ng-pass!", ip="10.0.0.9").status_code, 429)

    def test_success_resets_the_counter(self):
        for _ in range(4):
            self.login("wrong")
        self.assertEqual(self.login("Str0ng-pass!").status_code, 200)
        for _ in range(4):
            self.login("wrong")

        self.assertEqual(self.login("Str0ng-pass!").status_code, 200)

    def test_admin_login_shares_the_lockout(self):
        staff = User.objects.create_user(
            phone="+998901110009",
            password="Str0ng-pass!",
            shop=self.shop_a,
            full_name="Admin",
            is_staff=True,
            is_superuser=True,
        )
        web = Client(REMOTE_ADDR="10.1.1.1")
        for _ in range(5):
            web.post("/admin/login/", {"username": staff.phone, "password": "wrong"})

        web.post("/admin/login/", {"username": staff.phone, "password": "Str0ng-pass!"})

        self.assertNotIn("_auth_user_id", web.session)

    def test_a_spoofed_forwarded_for_header_does_not_dodge_the_limit(self):
        # No proxy configured (NUM_PROXIES=0): the header is ignored, so changing
        # it on every try gets the attacker nowhere.
        for n in range(5):
            APIClient(REMOTE_ADDR="10.0.0.1").post(
                "/api/auth/login/",
                {"phone": "+998901110001", "password": "wrong"},
                format="json",
                HTTP_X_FORWARDED_FOR=f"9.9.9.{n}",
            )

        response = APIClient(REMOTE_ADDR="10.0.0.1").post(
            "/api/auth/login/",
            {"phone": "+998901110001", "password": "Str0ng-pass!"},
            format="json",
            HTTP_X_FORWARDED_FOR="1.2.3.4",
        )

        self.assertEqual(response.status_code, 429)

    def test_behind_proxies_the_real_client_address_is_used(self):
        # Caddy + nginx: X-Forwarded-For is "client, caddy"; the client is 2nd from the right.
        headers = {"HTTP_X_FORWARDED_FOR": "5.5.5.5, 172.18.0.1", "REMOTE_ADDR": "172.18.0.2"}
        with override_settings(REST_FRAMEWORK={**self.rest_framework(), "NUM_PROXIES": 2}):
            for _ in range(5):
                APIClient().post(
                    "/api/auth/login/",
                    {"phone": "+998901110001", "password": "wrong"},
                    format="json",
                    **headers,
                )
            attacker = APIClient().post(
                "/api/auth/login/",
                {"phone": "+998901110001", "password": "Str0ng-pass!"},
                format="json",
                **headers,
            )
            other_client = APIClient().post(
                "/api/auth/login/",
                {"phone": "+998901110001", "password": "Str0ng-pass!"},
                format="json",
                HTTP_X_FORWARDED_FOR="8.8.8.8, 172.18.0.1",
                REMOTE_ADDR="172.18.0.2",
            )

        self.assertEqual(attacker.status_code, 429)
        self.assertEqual(other_client.status_code, 200)

    @staticmethod
    def rest_framework():
        from django.conf import settings

        return settings.REST_FRAMEWORK


class SessionRevocationTests(TwoShopsTestCase):
    def refresh_with(self, cookie_value):
        client = APIClient()
        client.cookies["refresh_token"] = cookie_value
        return client.post("/api/auth/refresh/", {}, format="json")

    def test_changing_the_password_kills_every_older_refresh_token(self):
        api = APIClient()
        login = api.post(
            "/api/auth/login/",
            {"phone": "+998901110001", "password": "Str0ng-pass!"},
            format="json",
        )
        old_cookie = login.cookies["refresh_token"].value
        access = login.json()["data"]["access"]
        self.assertEqual(self.refresh_with(old_cookie).status_code, 200)

        changed = api.post(
            "/api/auth/password/",
            {"old_password": "Str0ng-pass!", "new_password": "An0ther-pass!"},
            format="json",
            HTTP_AUTHORIZATION=f"Bearer {access}",
        )

        self.assertEqual(changed.status_code, 200)
        self.assertEqual(self.refresh_with(old_cookie).status_code, 401)  # the stolen one is dead
        self.assertEqual(self.refresh_with(changed.cookies["refresh_token"].value).status_code, 200)

    def test_logout_revokes_the_refresh_token(self):
        api = APIClient()
        login = api.post(
            "/api/auth/login/",
            {"phone": "+998901110001", "password": "Str0ng-pass!"},
            format="json",
        )
        cookie = login.cookies["refresh_token"].value
        api.post(
            "/api/auth/logout/",
            {},
            format="json",
            HTTP_AUTHORIZATION=f"Bearer {login.json()['data']['access']}",
        )

        self.assertEqual(self.refresh_with(cookie).status_code, 401)


class RateLimitTests(TwoShopsTestCase):
    def test_a_logged_in_user_is_rate_limited(self):
        with patch.object(UserRateThrottle, "THROTTLE_RATES", {"user": "3/min"}):
            codes = [self.client.get("/api/categories/").status_code for _ in range(5)]

        self.assertEqual(codes, [200, 200, 200, 429, 429])

    def test_report_files_have_their_own_tighter_limit(self):
        from rest_framework.throttling import ScopedRateThrottle

        with patch.object(ScopedRateThrottle, "THROTTLE_RATES", {"exports": "2/min"}):
            codes = [
                self.client.get(
                    "/api/reports/export/", {"period": "day", "file": "csv"}
                ).status_code
                for _ in range(3)
            ]

        self.assertEqual(codes, [200, 200, 429])


class ExportFormulaInjectionTests(TwoShopsTestCase):
    """A product named =HYPERLINK(...) must open in Excel as text, not run."""

    EVIL = '=HYPERLINK("http://evil.example","click")'

    def test_names_starting_with_a_formula_character_are_neutralised(self):
        product = Product.objects.create(
            shop=self.shop_a,
            name=self.EVIL,
            unit=Product.Unit.PIECE,
            markup_amount=2000,
            min_stock=Decimal("50"),
        )
        Batch.objects.create(
            shop=self.shop_a,
            product=product,
            qty_initial=10,
            qty_remaining=10,
            cost_price=1000,
            sale_price=1200,
        )
        create_sale(
            shop=self.shop_a,
            user=self.owner_a,
            client_id=uuid4(),
            payment_type=Sale.PaymentType.CASH,
            customer=None,
            cart=[CartLine(product=product, qty=Decimal("1"))],
        )

        month = self.client.get("/api/reports/export/", {"period": "month", "file": "xlsx"})
        low = self.client.get("/api/reports/low-stock/export/")

        for response in (month, low):
            workbook = load_workbook(BytesIO(response.content))
            for sheet in workbook:
                for row in sheet.iter_rows():
                    for cell in row:
                        if isinstance(cell.value, str):
                            self.assertFalse(
                                cell.value.startswith(("=", "+", "-", "@")),
                                f"live formula in {sheet.title}!{cell.coordinate}: {cell.value}",
                            )
        products_sheet = load_workbook(BytesIO(month.content))["Mahsulotlar bo'yicha"]
        self.assertEqual(products_sheet["A2"].value, "'" + self.EVIL)  # still readable as text


class ImageUploadTests(TwoShopsTestCase):
    def upload(self, name, content, content_type):
        with override_settings(MEDIA_ROOT=tempfile.mkdtemp()):
            return self.client.post(
                "/api/products/",
                {
                    "name": "Foto",
                    "unit": "piece",
                    "markup_amount": "1000",
                    "min_stock": "0",
                    "image": SimpleUploadedFile(name, content, content_type=content_type),
                },
                format="multipart",
            )

    @staticmethod
    def picture(fmt, size=(20, 20)):
        buffer = BytesIO()
        Image.new("RGB", size, "red").save(buffer, fmt)
        return buffer.getvalue()

    def test_a_real_png_and_jpeg_are_accepted(self):
        self.assertEqual(self.upload("a.png", self.picture("PNG"), "image/png").status_code, 201)
        self.assertEqual(self.upload("a.jpg", self.picture("JPEG"), "image/jpeg").status_code, 201)

    def test_scripts_html_and_svg_pretending_to_be_pictures_are_rejected(self):
        for name, content, kind in [
            ("shell.jpg", b"<?php system($_GET['c']); ?>", "image/jpeg"),
            ("page.png", b"<html><script>alert(1)</script></html>", "image/png"),
            (
                "logo.svg",
                b'<svg xmlns="http://www.w3.org/2000/svg"><script>alert(1)</script></svg>',
                "image/svg+xml",
            ),
        ]:
            with self.subTest(name=name):
                self.assertEqual(self.upload(name, content, kind).status_code, 400)

    def test_other_picture_formats_are_rejected(self):
        self.assertEqual(self.upload("a.gif", self.picture("GIF"), "image/gif").status_code, 400)

    def test_oversized_uploads_are_rejected(self):
        big = BytesIO()
        Image.frombytes("RGB", (2000, 1100), os.urandom(2000 * 1100 * 3)).save(big, "PNG")
        self.assertGreater(len(big.getvalue()), 5 * 1024 * 1024)

        self.assertEqual(self.upload("big.png", big.getvalue(), "image/png").status_code, 400)


class DocsAndProductionConfigTests(APITestCase):
    def test_api_docs_are_not_published_when_debug_is_off(self):
        self.assertEqual(self.client.get("/api/schema/").status_code, 404)
        self.assertEqual(self.client.get("/api/schema/swagger-ui/").status_code, 404)

    def test_production_settings_are_locked_down(self):
        """Loads config.settings.prod in a clean interpreter and checks the knobs
        that matter, so a careless edit can't quietly open them up."""
        script = """
# Ignore the developer's own .env: only the environment, so the code's defaults show.
import decouple
decouple.config = decouple.Config(decouple.RepositoryEmpty())
import django
django.setup()
from django.conf import settings as s
from django.urls import NoReverseMatch, reverse
import config.urls as urls

try:
    reverse("admin:index")
    admin = True
except NoReverseMatch:
    admin = False
docs = any(getattr(u, "name", None) == "schema" for u in urls.urlpatterns)
print(admin, s.DEBUG, s.CORS_ALLOWED_ORIGINS, s.SESSION_COOKIE_SECURE, s.CSRF_COOKIE_SECURE,
      s.SECURE_HSTS_SECONDS > 0, s.SECURE_SSL_REDIRECT, s.SESSION_COOKIE_SAMESITE, docs)
"""
        env = {
            **os.environ,
            "DJANGO_SETTINGS_MODULE": "config.settings.prod",
            "SECRET_KEY": "x" * 60,
            "DATABASE_URL": "postgresql://u:p@localhost:5432/db",
            "ALLOWED_HOSTS": "shop.example.uz",
        }
        for var in ("ADMIN_ENABLED", "CORS_ALLOWED_ORIGINS", "SECURE_SSL_REDIRECT", "NUM_PROXIES"):
            env.pop(var, None)

        result = subprocess.run(
            [sys.executable, "-c", script], capture_output=True, text=True, env=env, timeout=60
        )

        self.assertEqual(result.returncode, 0, result.stderr)
        # admin off, debug off, no CORS origins, secure cookies, HSTS, https redirect,
        # strict same-site, no API docs route
        self.assertEqual(
            result.stdout.split("\n")[-2].strip(), "False False [] True True True True Strict False"
        )

    def test_plain_http_test_mode_exists_but_only_when_asked_for(self):
        """ALLOW_INSECURE_HTTP relaxes exactly the TLS-only settings (no-domain
        test servers) and is off unless the env var says so."""
        script = """
import decouple
decouple.config = decouple.Config(decouple.RepositoryEmpty())
import django
django.setup()
from django.conf import settings as s
print(s.SECURE_SSL_REDIRECT, s.SESSION_COOKIE_SECURE, s.REFRESH_COOKIE_SECURE, s.SECURE_HSTS_SECONDS)
"""
        base_env = {
            **os.environ,
            "DJANGO_SETTINGS_MODULE": "config.settings.prod",
            "SECRET_KEY": "x" * 60,
            "DATABASE_URL": "postgresql://u:p@localhost:5432/db",
            "ALLOWED_HOSTS": "203-0-113-5.sslip.io",
        }
        base_env.pop("ALLOW_INSECURE_HTTP", None)

        def run(extra):
            out = subprocess.run(
                [sys.executable, "-c", script],
                capture_output=True,
                text=True,
                env={**base_env, **extra},
                timeout=60,
            )
            self.assertEqual(out.returncode, 0, out.stderr)
            return out.stdout.strip().split("\n")[-1]

        self.assertEqual(run({}), "True True True 31536000")
        self.assertEqual(run({"ALLOW_INSECURE_HTTP": "True"}), "False False False 0")
